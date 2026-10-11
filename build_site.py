#!/usr/bin/env python3
"""Build docs/index.html from the training-plan workbook.

Usage:  python3 build_site.py
Reads   富士山馬拉松_2026_課表.xlsx (sheets 課表 / 週總覽)
Writes  docs/index.html  (single self-contained page, phone-first)

The page only renders what is in the workbook. The verdict for each run is
typed by hand into column V (判定) and column W (提醒, items separated by ｜);
nothing is judged automatically here.

Run this every time the xlsx changes and commit both files together.
"""
from __future__ import annotations

import datetime as dt
import json
import pathlib

import openpyxl

ROOT = pathlib.Path(__file__).resolve().parent
XLSX = ROOT / "富士山馬拉松_2026_課表.xlsx"
OUT = ROOT / "docs" / "index.html"
RACE_DAY = dt.date(2026, 12, 14)
JST = dt.timezone(dt.timedelta(hours=9))

# 課表 columns: A 週 B 日期 C 星期 D 階段 E 類型 F 計畫km G 配速/心拍 H 內容 I 備註
# J 實際km K 完成 L 心得 M 時間 N 平均配速 O 平均心拍 P 最大心拍 Q 平均功率 R 卡路里 S 天氣 T 鞋 U 裝置
# V 判定 W 提醒
COLS = [
    "week", "date", "dow", "phase", "type", "plan_km", "target", "content", "note",
    "actual_km", "done", "memo", "time", "pace", "hr", "hr_max", "power", "kcal",
    "weather", "shoes", "device", "verdict", "reminder",
]
VERDICTS = ("達成", "達成但有提醒", "未達成", "跳過")


def num(v):
    if v is None or v == "":
        return None
    try:
        return float(v)
    except (TypeError, ValueError):
        return None


def text(v):
    if v is None:
        return None
    s = str(v).strip()
    return s or None


def load_rows() -> list[dict]:
    wb = openpyxl.load_workbook(XLSX, data_only=True)
    ws = wb["課表"]
    rows = []
    for r in ws.iter_rows(min_row=2, values_only=True):
        if not isinstance(r[1], dt.datetime):
            continue
        r = tuple(r) + (None,) * (len(COLS) - len(r))
        row = dict(zip(COLS, r[: len(COLS)]))
        row["date"] = row["date"].date().isoformat()
        for k in ("plan_km", "actual_km", "hr", "hr_max", "power", "kcal", "week"):
            row[k] = num(row[k])
        row["week"] = int(row["week"]) if row["week"] is not None else None
        for k in ("time", "pace", "memo", "note", "target", "content", "weather", "shoes", "device"):
            row[k] = text(row[k])
        row["done"] = bool(row["done"]) or row["actual_km"] is not None
        row["is_rest"] = (row["plan_km"] or 0) == 0
        v = text(row["verdict"])
        if v and v not in VERDICTS:
            print(f"  warning: {row['date']} 判定 '{v}' is not one of {VERDICTS}")
        row["verdict"] = v
        w = text(row["reminder"])
        row["reminders"] = [s.strip() for s in w.split("｜") if s.strip()] if w else []
        del row["reminder"]
        rows.append(row)
    rows.sort(key=lambda x: x["date"])
    return rows


def load_weeks(rows: list[dict]) -> list[dict]:
    wb = openpyxl.load_workbook(XLSX, data_only=True)
    ws = wb["週總覽"]
    focus = {}
    for r in ws.iter_rows(min_row=2, values_only=True):
        if isinstance(r[0], (int, float)):
            focus[int(r[0])] = r[8]
    weeks: dict[int, dict] = {}
    for row in rows:
        w = row["week"]
        if w is None:
            continue
        wk = weeks.setdefault(w, {"week": w, "start": row["date"], "end": row["date"], "phase": row["phase"],
                                  "plan_km": 0.0, "actual_km": 0.0, "long_km": 0.0, "focus": focus.get(w)})
        wk["end"] = row["date"]
        wk["plan_km"] += row["plan_km"] or 0
        wk["actual_km"] += row["actual_km"] or 0
        wk["long_km"] = max(wk["long_km"], row["plan_km"] or 0)
    for wk in weeks.values():
        wk["plan_km"] = round(wk["plan_km"], 3)
        wk["actual_km"] = round(wk["actual_km"], 2)
    return [weeks[k] for k in sorted(weeks)]


TEMPLATE = r"""<!doctype html>
<html lang="zh-Hant">
<head>
<meta charset="utf-8">
<meta name="viewport" content="width=device-width, initial-scale=1, viewport-fit=cover">
<meta name="color-scheme" content="light dark">
<title>富士山マラソン 2026</title>
<meta name="description" content="富士山マラソン 2026 訓練紀錄：每日跑步報告、下次練習、歷史紀錄">
<style>
:root{
  --bg:#fff; --surface:#fff; --ink:#0a0a0a; --ink-2:#555; --ink-3:#999;
  --line:#e6e6e6; --chip:#f2f2f2; --accent:#1d6fe0; --accent-soft:#e3edfb;
  --good:#1a8f4e; --good-soft:#e2f4ea; --warn:#c7700a; --warn-soft:#fbeedc;
  --bad:#c62828; --bad-soft:#fbe3e3; --skip:#6b6b6b; --skip-soft:#ededed;
  --track:#ececec;
}
@media (prefers-color-scheme: dark){
  :root{
    --bg:#000; --surface:#000; --ink:#f3f3f3; --ink-2:#a8a8a8; --ink-3:#6f6f6f;
    --line:#262626; --chip:#161616; --accent:#6ea2ff; --accent-soft:#132544;
    --good:#4ccb84; --good-soft:#0f2a1b; --warn:#f0a44a; --warn-soft:#2e1f0a;
    --bad:#ff7b7b; --bad-soft:#3a1212; --skip:#a3a3a3; --skip-soft:#1c1c1c;
    --track:#1f1f1f;
  }
}
*{box-sizing:border-box}
html,body{margin:0;background:var(--bg);color:var(--ink);
  font:15px/1.55 -apple-system,BlinkMacSystemFont,"Helvetica Neue","Hiragino Sans","Noto Sans CJK TC","PingFang TC",sans-serif;
  -webkit-font-smoothing:antialiased}
a{color:inherit}
.wrap{max-width:600px;margin:0 auto;padding:0 16px env(safe-area-inset-bottom)}
header{position:sticky;top:0;z-index:2;background:color-mix(in srgb,var(--bg) 88%,transparent);
  backdrop-filter:saturate(180%) blur(12px);-webkit-backdrop-filter:saturate(180%) blur(12px);
  border-bottom:1px solid var(--line)}
header .wrap{display:flex;align-items:center;justify-content:space-between;height:52px}
header h1{font-size:17px;font-weight:700;margin:0;letter-spacing:-.01em}
header .count{font-size:13px;color:var(--ink-2);font-variant-numeric:tabular-nums}
.hero{padding:20px 0 16px;border-bottom:1px solid var(--line)}
.hero .big{display:flex;align-items:baseline;gap:6px}
.hero .big b{font-size:34px;font-weight:800;letter-spacing:-.02em;font-variant-numeric:tabular-nums}
.hero .big span{color:var(--ink-2)}
.hero .sub{color:var(--ink-2);font-size:13px;margin-top:2px}
.bar{height:6px;background:var(--track);border-radius:4px;overflow:hidden;margin-top:10px}
.bar i{display:block;height:100%;background:var(--accent);border-radius:4px;width:0;transition:width .6s ease}
.bar.done i{background:var(--good)}
h2{font-size:13px;font-weight:600;color:var(--ink-3);text-transform:uppercase;letter-spacing:.06em;margin:22px 0 6px}
.post{display:grid;grid-template-columns:40px 1fr;gap:12px;padding:14px 0;border-bottom:1px solid var(--line)}
.avatar{width:36px;height:36px;border-radius:50%;background:var(--chip);display:flex;align-items:center;justify-content:center;font-size:18px}
.avatar.hi{background:var(--accent-soft)}
.meta{display:flex;align-items:center;gap:6px;flex-wrap:wrap;font-size:14px}
.meta b{font-weight:600}
.meta .dim{color:var(--ink-3)}
.tag{font-size:11px;font-weight:600;padding:2px 7px;border-radius:999px;background:var(--chip);color:var(--ink-2)}
.tag.accent{background:var(--accent-soft);color:var(--accent)}
.tag.good{background:var(--good-soft);color:var(--good)}
.tag.warn{background:var(--warn-soft);color:var(--warn)}
.tag.bad{background:var(--bad-soft);color:var(--bad)}
.tag.skip{background:var(--skip-soft);color:var(--skip)}
.body{margin-top:4px}
.body p{margin:4px 0}
.title{font-size:16px;font-weight:600;margin:2px 0 4px}
.muted{color:var(--ink-2)}
.small{font-size:13px}
.stats{display:grid;grid-template-columns:repeat(3,1fr);gap:8px;margin:10px 0 6px}
.stat{background:var(--chip);border-radius:10px;padding:8px 10px}
.stat b{display:block;font-size:17px;font-weight:700;letter-spacing:-.01em;font-variant-numeric:tabular-nums}
.stat span{font-size:11px;color:var(--ink-3)}
.verdict{display:flex;gap:8px;align-items:flex-start;margin-top:10px;padding:10px 12px;border-radius:12px;background:var(--chip);font-size:14px}
.verdict.good{background:var(--good-soft)} .verdict.warn{background:var(--warn-soft)}
.verdict.bad{background:var(--bad-soft)} .verdict.skip{background:var(--skip-soft)}
.verdict .dot{flex:none;width:8px;height:8px;border-radius:50%;margin-top:7px;background:var(--ink-3)}
.verdict.good .dot{background:var(--good)} .verdict.warn .dot{background:var(--warn)}
.verdict.bad .dot{background:var(--bad)} .verdict.skip .dot{background:var(--skip)}
ul.points{margin:8px 0 0;padding-left:0;list-style:none;color:var(--ink-2);font-size:14px}
ul.points li{padding-left:14px;position:relative;margin:3px 0}
ul.points li::before{content:"";position:absolute;left:0;top:9px;width:5px;height:5px;border-radius:50%;background:var(--ink-3)}
.memo{margin-top:8px;font-size:15px}
.foot{color:var(--ink-3);font-size:12px;margin-top:8px}
.week{padding:12px 0;border-bottom:1px solid var(--line)}
.week .row{display:flex;justify-content:space-between;align-items:baseline;gap:8px}
.week .row b{font-weight:600}
.week .row .num{font-variant-numeric:tabular-nums;color:var(--ink-2);font-size:13px}
.week .focus{color:var(--ink-3);font-size:12px;margin-top:2px}
.week.now .row b::after{content:"本週";margin-left:6px;font-size:11px;font-weight:600;color:var(--accent);background:var(--accent-soft);padding:1px 6px;border-radius:999px}
.week.future{opacity:.6}
.empty{color:var(--ink-3);padding:14px 0;border-bottom:1px solid var(--line)}
footer{padding:28px 0 40px;color:var(--ink-3);font-size:12px;text-align:center}
footer a{color:var(--ink-2)}
</style>
</head>
<body>
<header><div class="wrap"><h1>🗻 富士山マラソン 2026</h1><div class="count" id="countdown"></div></div></header>
<main class="wrap">
  <section class="hero" id="hero"></section>

  <h2>今天</h2>
  <div id="today"></div>

  <h2>下次練習</h2>
  <div id="next" class="next-list"></div>

  <h2>最新報告</h2>
  <div id="latest"></div>

  <h2>週進度</h2>
  <div id="weeks"></div>

  <h2>歷史紀錄</h2>
  <div id="history"></div>
</main>
<footer>
  資料更新：<span id="built"></span><br>
  <a href="https://github.com/iamwuwei/running" target="_blank" rel="noopener">GitHub</a> ·
  <a href="https://github.com/iamwuwei/running/raw/main/%E5%AF%8C%E5%A3%AB%E5%B1%B1%E9%A6%AC%E6%8B%89%E6%9D%BE_2026_%E8%AA%B2%E8%A1%A8.xlsx">下載 Excel</a>
</footer>
<script id="data" type="application/json">__DATA__</script>
<script>
const D = JSON.parse(document.getElementById('data').textContent);
const esc = s => String(s ?? '').replace(/[&<>"']/g, c => ({'&':'&amp;','<':'&lt;','>':'&gt;','"':'&quot;',"'":'&#39;'}[c]));
const fmtKm = v => v == null ? '–' : (Math.round(v*100)/100).toString();
// "today" in Japan, independent of the viewer's timezone
const jst = new Date(Date.now() + 9*3600*1000);
const today = jst.toISOString().slice(0,10);
const race = D.race_day;
const dayDiff = (a,b) => Math.round((Date.parse(b) - Date.parse(a)) / 86400000);
const addDays = (iso, n) => new Date(Date.parse(iso) + n*86400000).toISOString().slice(0,10);
const md = iso => `${+iso.slice(5,7)}/${+iso.slice(8,10)}`;
const dowName = r => r.dow ? `週${r.dow}` : '';
const typeTag = r => r.is_rest ? '' : `<span class="tag accent">${esc(r.type)}</span>`;
const avatarFor = r => r.type === '肌力' ? '💪' : r.is_rest ? '😴' : r.type === '比賽' ? '🏁' : '🏃';
// verdict text (column V) -> colour class; anything else renders neutral
const VERDICT_CLASS = {'達成':'good', '達成但有提醒':'warn', '未達成':'bad', '跳過':'skip'};

// ---- header / hero
const toRace = dayDiff(today, race);
document.getElementById('countdown').textContent =
  toRace > 0 ? `比賽倒數 ${toRace} 天` : toRace === 0 ? '比賽日 🏁' : '比賽結束';

const done = D.rows.filter(r => r.actual_km != null);
const totalActual = done.reduce((s,r) => s + r.actual_km, 0);
const planToDate = D.rows.filter(r => r.date <= today).reduce((s,r) => s + (r.plan_km||0), 0);
const planTotal = D.rows.reduce((s,r) => s + (r.plan_km||0), 0);
const pct = planToDate ? totalActual/planToDate*100 : 0;
document.getElementById('hero').innerHTML = `
  <div class="big"><b>${fmtKm(totalActual)}</b><span>km 累積</span></div>
  <div class="sub">課表至今 ${fmtKm(planToDate)} km · 達成 ${Math.round(pct)}% · ${done.length} 次練習 · 全程計畫 ${fmtKm(planTotal)} km</div>
  <div class="bar ${pct>=100?'done':''}"><i style="width:${Math.min(100,pct)}%"></i></div>`;

// ---- plan card (today / next)
function planCard(r, {hi=false, label=''} = {}) {
  const when = r.date === today ? '今天' : r.date === addDays(today,1) ? '明天' : `${md(r.date)} ${dowName(r)}`;
  const body = r.is_rest
    ? `<div class="title">${esc(r.content || '休息')}</div>${r.note ? `<p class="muted">${esc(r.note)}</p>` : ''}`
    : `<div class="title">${esc(r.content)}</div>
       <p class="muted">${esc(r.target || '')}</p>
       ${r.note ? `<p class="muted small">📝 ${esc(r.note)}</p>` : ''}`;
  return `<article class="post">
    <div class="avatar ${hi?'hi':''}">${avatarFor(r)}</div>
    <div><div class="meta"><b>${esc(when)}</b>${label ? `<span class="dim">· ${esc(label)}</span>` : ''} ${typeTag(r)} ${r.plan_km ? `<span class="tag">${fmtKm(r.plan_km)} km</span>` : ''}</div>
    <div class="body">${body}<div class="foot">第 ${r.week} 週 · ${esc(r.phase)}</div></div></div></article>`;
}

// ---- report card (completed run): numbers from the sheet, verdict from V/W verbatim
function reportCard(r) {
  const vcls = r.verdict ? (VERDICT_CLASS[r.verdict] || '') : '';
  const stats = [
    [`課表 ${fmtKm(r.plan_km)} km`, r.actual_km != null ? `${fmtKm(r.actual_km)} km` : null, null],
    ['時間', r.time, null],
    ['配速', r.pace ? `${r.pace}/km` : null, null],
    ['心拍', r.hr != null ? `${Math.round(r.hr)}` : null, r.hr_max ? `最高 ${Math.round(r.hr_max)}` : null],
    ['卡路里', r.kcal != null ? `${Math.round(r.kcal)}` : null, null],
    ['功率', r.power != null ? `${Math.round(r.power)} W` : null, null],
  ].filter(s => s[1]);
  const verdictBlock = r.verdict ? `<div class="verdict ${vcls}"><span class="dot"></span><div><b>${esc(r.verdict)}</b>
      ${r.reminders.length ? `<ul class="points">${r.reminders.map(p => `<li>${esc(p)}</li>`).join('')}</ul>` : ''}</div></div>` : '';
  return `<article class="post">
    <div class="avatar">${avatarFor(r)}</div>
    <div>
      <div class="meta"><b>${md(r.date)} ${dowName(r)}</b><span class="dim">· 第 ${r.week} 週</span> ${typeTag(r)}
        ${r.verdict ? `<span class="tag ${vcls}">${esc(r.verdict)}</span>` : ''}</div>
      <div class="body">
        <div class="title">${esc(r.content)}</div>
        ${r.memo ? `<p class="memo">${esc(r.memo)}</p>` : ''}
        <div class="stats">${stats.map(([k,v,sub]) => `<div class="stat"><b>${esc(v)}</b><span>${esc(k)}${sub ? ' · ' + esc(sub) : ''}</span></div>`).join('')}</div>
        ${verdictBlock}
        <div class="foot">${[r.weather, r.shoes, r.device].filter(Boolean).map(esc).join(' · ')}</div>
      </div>
    </div></article>`;
}

// ---- today
const todayRow = D.rows.find(r => r.date === today);
document.getElementById('today').innerHTML = todayRow
  ? (todayRow.actual_km != null ? reportCard(todayRow) : planCard(todayRow, {hi:true}))
  : `<div class="empty">${today > race ? '課表已結束，辛苦了 🎉' : '課表還沒開始'}</div>`;

// ---- next workouts: upcoming runs after today (rest days skipped), first 3
const upcoming = D.rows.filter(r => r.date > today && !r.is_rest).slice(0,3);
document.getElementById('next').innerHTML = upcoming.length
  ? upcoming.map((r,i) => planCard(r, {hi:i===0, label: `${dayDiff(today, r.date)} 天後`})).join('')
  : `<div class="empty">沒有後續課表</div>`;

// ---- latest report = newest run not already shown as today; history = the rest
const shownToday = todayRow && todayRow.actual_km != null ? todayRow : null;
const rest = done.filter(r => r !== shownToday);
const latest = rest.length ? rest[rest.length-1] : null;
document.getElementById('latest').innerHTML = latest ? reportCard(latest)
  : `<div class="empty">${shownToday ? '今天的就是第一筆紀錄' : '還沒有紀錄'}</div>`;

// ---- weeks
document.getElementById('weeks').innerHTML = D.weeks.map(w => {
  const p = w.plan_km ? Math.min(100, w.actual_km / w.plan_km * 100) : 0;
  const now = today >= w.start && today <= w.end;
  const future = today < w.start;
  const complete = w.actual_km >= w.plan_km && w.plan_km > 0;
  return `<div class="week ${now?'now':''} ${future?'future':''}">
    <div class="row"><b>第 ${w.week} 週</b><span class="num">${fmtKm(w.actual_km)} / ${fmtKm(w.plan_km)} km · ${Math.round(p)}%</span></div>
    <div class="focus">${md(w.start)}–${md(w.end)} · ${esc(w.phase)}${w.focus ? ' · ' + esc(w.focus) : ''}</div>
    <div class="bar ${complete?'done':''}"><i style="width:${p}%"></i></div></div>`;
}).join('');

// ---- history (newest first)
const hist = rest.filter(r => r !== latest).reverse();
document.getElementById('history').innerHTML = hist.length ? hist.map(reportCard).join('')
  : `<div class="empty">更早的紀錄會出現在這裡</div>`;

document.getElementById('built').textContent = D.built;
</script>
</body>
</html>
"""


def build() -> None:
    rows = load_rows()
    weeks = load_weeks(rows)
    data = {
        "race_day": RACE_DAY.isoformat(),
        "built": dt.datetime.now(JST).strftime("%Y/%m/%d %H:%M JST"),
        "rows": rows,
        "weeks": weeks,
    }
    payload = json.dumps(data, ensure_ascii=False).replace("</", "<\\/")
    OUT.parent.mkdir(parents=True, exist_ok=True)
    OUT.write_text(TEMPLATE.replace("__DATA__", payload), encoding="utf-8")
    done = [r for r in rows if r["actual_km"] is not None]
    print(f"wrote {OUT.relative_to(ROOT)}  rows={len(rows)} completed={len(done)} weeks={len(weeks)}")
    for r in done:
        print(f"  {r['date']} {r['type']:<8} {r['actual_km']:>5} km  判定={r['verdict'] or '（未評）'}  提醒={len(r['reminders'])} 條")


if __name__ == "__main__":
    build()
